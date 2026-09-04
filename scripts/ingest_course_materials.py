#!/usr/bin/env python3
"""Build a tutor-friendly evidence bundle from course materials.

The script extracts structure and text; the tutor remains responsible for
visually inspecting formula-, diagram-, and screenshot-heavy pages.
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import os
import re
import shutil
import sys
import zipfile
from pathlib import Path
from typing import Any


SUPPORTED = {".pptx", ".docx", ".pdf", ".png", ".jpg", ".jpeg", ".webp"}
DEPENDENCIES = {
    ".pptx": ("pptx", "python-pptx"),
    ".docx": ("docx", "python-docx"),
    ".pdf": ("pdfplumber", "pdfplumber"),
}


def safe_name(value: str) -> str:
    cleaned = re.sub(r"[^\w.-]+", "_", value, flags=re.UNICODE).strip("._")
    return cleaned or "source"


def table_rows(table: Any) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def extract_pptx(path: Path, image_dir: Path) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(str(path))
    slides = []
    images = []
    for number, slide in enumerate(deck.slides, 1):
        title = ""
        blocks = []
        tables = []
        for shape_index, shape in enumerate(slide.shapes, 1):
            if getattr(shape, "has_table", False):
                tables.append(table_rows(shape.table))
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                )
                if text:
                    if getattr(shape, "is_placeholder", False) and not title:
                        title = text.splitlines()[0]
                    blocks.append(text)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                extension = shape.image.ext or "bin"
                destination = image_dir / f"slide{number}_image{shape_index}.{extension}"
                destination.write_bytes(shape.image.blob)
                images.append({
                    "path": str(destination),
                    "location": {"slide": number},
                    "context": title or "\n".join(blocks)[:200],
                })
        notes = ""
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except (AttributeError, KeyError):
            pass
        slides.append({
            "number": number,
            "title": title,
            "text": "\n".join(blocks),
            "tables": tables,
            "notes": notes,
        })
    return {"kind": "pptx", "units": slides, "images": images}


def extract_docx(path: Path, image_dir: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(str(path))
    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            paragraphs.append({
                "number": index,
                "style": paragraph.style.name if paragraph.style else "",
                "text": text,
            })
    tables = [table_rows(table) for table in document.tables]
    images = []
    with zipfile.ZipFile(path) as archive:
        for member in archive.namelist():
            if not member.startswith("word/media/") or member.endswith("/"):
                continue
            destination = image_dir / safe_name(Path(member).name)
            destination.write_bytes(archive.read(member))
            images.append({"path": str(destination), "location": {"part": member}})
    return {
        "kind": "docx",
        "units": paragraphs,
        "tables": tables,
        "images": images,
    }


def extract_pdf(path: Path, image_dir: Path) -> dict[str, Any]:
    import pdfplumber

    pages = []
    with pdfplumber.open(str(path)) as document:
        for number, page in enumerate(document.pages, 1):
            pages.append({
                "number": number,
                "text": (page.extract_text() or "").strip(),
                "tables": page.extract_tables() or [],
                "width": page.width,
                "height": page.height,
            })
    images = []
    if importlib.util.find_spec("fitz"):
        import fitz

        document = fitz.open(str(path))
        try:
            for page_index in range(len(document)):
                page = document[page_index]
                for image_index, info in enumerate(page.get_images(full=True), 1):
                    payload = document.extract_image(info[0])
                    destination = image_dir / (
                        f"page{page_index + 1}_image{image_index}.{payload['ext']}"
                    )
                    destination.write_bytes(payload["image"])
                    images.append({
                        "path": str(destination),
                        "location": {"page": page_index + 1},
                    })
        finally:
            document.close()
    return {"kind": "pdf", "units": pages, "images": images}


def extract_image(path: Path, _image_dir: Path) -> dict[str, Any]:
    return {
        "kind": "image",
        "units": [],
        "images": [{"path": str(path), "location": {"source": "standalone"}}],
        "visual_review_required": True,
    }


def missing_dependency(extension: str) -> str | None:
    requirement = DEPENDENCIES.get(extension)
    if not requirement:
        return None
    module, package = requirement
    return None if importlib.util.find_spec(module) else package


def discover(target: Path) -> list[Path]:
    if target.is_file():
        return [target] if target.suffix.lower() in SUPPORTED else []
    ignored = {"_tutor_materials", ".git", "__pycache__"}
    return sorted(
        path for path in target.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED
        and not any(part in ignored for part in path.parts)
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("target", nargs="?", help="course file or directory")
    parser.add_argument("--output", help="output directory (default: TARGET/_tutor_materials)")
    parser.add_argument("--check", action="store_true", help="report optional dependency status")
    args = parser.parse_args()

    if args.check:
        for extension, (module, package) in DEPENDENCIES.items():
            state = "available" if importlib.util.find_spec(module) else "missing"
            print(f"{extension}: {state} ({package})")
        print("PDF embedded-image extraction:", "available" if importlib.util.find_spec("fitz") else "optional PyMuPDF missing")
        return 0
    if not args.target:
        parser.error("target is required unless --check is used")

    target = Path(args.target).expanduser().resolve()
    if not target.exists():
        print(f"Target not found: {target}", file=sys.stderr)
        return 2
    output = Path(args.output).expanduser().resolve() if args.output else (
        (target if target.is_dir() else target.parent) / "_tutor_materials"
    )
    output.mkdir(parents=True, exist_ok=True)

    sources = []
    warnings = []
    extractors = {
        ".pptx": extract_pptx,
        ".docx": extract_docx,
        ".pdf": extract_pdf,
        ".png": extract_image,
        ".jpg": extract_image,
        ".jpeg": extract_image,
        ".webp": extract_image,
    }
    for path in discover(target):
        extension = path.suffix.lower()
        missing = missing_dependency(extension)
        if missing:
            warnings.append(f"Skipped {path.name}: optional dependency {missing} is unavailable")
            continue
        image_dir = output / "images" / safe_name(path.stem)
        image_dir.mkdir(parents=True, exist_ok=True)
        try:
            result = extractors[extension](path, image_dir)
            result.update({"source": str(path), "filename": path.name})
            for image in result.get("images", []):
                image["path"] = os.path.relpath(image["path"], output)
            sources.append(result)
        except Exception as error:  # keep partial bundles useful
            warnings.append(f"Failed {path.name}: {type(error).__name__}: {error}")
            shutil.rmtree(image_dir, ignore_errors=True)

    bundle = {
        "schema_version": 1,
        "target": str(target),
        "source_count": len(sources),
        "sources": sources,
        "warnings": warnings,
        "tutor_instructions": {
            "evidence_first": "Distinguish extracted source content from tutor-added explanation.",
            "visual_review": "Inspect formula-, chart-, diagram-, and screenshot-heavy pages or images before teaching from them.",
            "diagnosis": "Map coverage, prerequisites, likely gaps, and the smallest next teaching step before generating a long lesson.",
        },
    }
    bundle_path = output / "_tutor_material_bundle.json"
    bundle_path.write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")
    print(bundle_path)
    for warning in warnings:
        print(f"WARNING: {warning}", file=sys.stderr)
    return 0 if sources else 1


if __name__ == "__main__":
    raise SystemExit(main())
